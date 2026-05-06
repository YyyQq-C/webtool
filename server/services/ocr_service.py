"""
OCR服务模块 - 图片转Excel/Word
支持多种OCR引擎
"""
import os
import io
import asyncio
from pathlib import Path
from typing import List, Dict, Any, Optional
from PIL import Image

# OCR引擎选择
OCR_ENGINE = os.environ.get("OCR_ENGINE", "paddleocr")  # paddleocr, tesseract, easyocr


def check_ocr_dependencies():
    """检查OCR依赖是否安装"""
    global OCR_ENGINE
    
    # 先检查 Tesseract（系统命令）
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        print("[OCR] Tesseract OCR 可用")
        OCR_ENGINE = "tesseract"
        return True
    except Exception as e:
        print(f"[OCR] Tesseract OCR 不可用: {e}")
    
    # 再检查 EasyOCR
    try:
        import easyocr
        print("[OCR] EasyOCR 可用")
        OCR_ENGINE = "easyocr"
        return True
    except Exception as e:
        print(f"[OCR] EasyOCR 不可用: {e}")
    
    # 最后尝试 PaddleOCR
    try:
        import os
        os.environ['FLAGS_use_mkldnn'] = '0'
        os.environ['FLAGS_enable_onednn'] = '0'
        from paddleocr import PaddleOCR
        print("[OCR] PaddleOCR 可用")
        OCR_ENGINE = "paddleocr"
        return True
    except Exception as e:
        print(f"[OCR] PaddleOCR 不可用: {e}")
    
    print("[OCR] 警告：没有可用的OCR引擎")
    return False


OCR_AVAILABLE = check_ocr_dependencies()


def ocr_image_tesseract(image_path: str, lang: str = 'chi_sim+eng') -> List[Dict]:
    """使用Tesseract OCR识别图片"""
    import pytesseract
    
    img = Image.open(image_path)
    
    # 获取文字和位置信息
    data = pytesseract.image_to_data(img, lang=lang, output_type=pytesseract.Output.DICT)
    
    results = []
    n_boxes = len(data['text'])
    
    for i in range(n_boxes):
        if int(data['conf'][i]) > 30:  # 置信度阈值
            text = data['text'][i].strip()
            if text:
                results.append({
                    'text': text,
                    'x': data['left'][i],
                    'y': data['top'][i],
                    'width': data['width'][i],
                    'height': data['height'][i],
                    'confidence': float(data['conf'][i]),
                })
    
    return results


def ocr_image_paddleocr(image_path: str) -> List[Dict]:
    """使用PaddleOCR识别图片"""
    # 禁用 ONEDNN/MKLDNN 避免兼容性问题
    import os
    os.environ['FLAGS_use_mkldnn'] = '0'
    os.environ['FLAGS_enable_onednn'] = '0'
    os.environ['FLAGS_enable_ir_optim'] = '0'  # 禁用 IR 优化
    
    from paddleocr import PaddleOCR
    
    # 使用 enable_mkldnn=False 避免问题
    ocr = PaddleOCR(use_angle_cls=True, lang='ch', enable_mkldnn=False)
    result = ocr.ocr(image_path)
    
    results = []
    if result and result[0]:
        for line in result[0]:
            box = line[0]  # [[x1,y1],[x2,y2],[x3,y3],[x4,y4]]
            text_info = line[1]  # (text, confidence)
            
            # 计算边界框
            x_coords = [p[0] for p in box]
            y_coords = [p[1] for p in box]
            x_min, x_max = min(x_coords), max(x_coords)
            y_min, y_max = min(y_coords), max(y_coords)
            
            results.append({
                'text': text_info[0],
                'x': int(x_min),
                'y': int(y_min),
                'width': int(x_max - x_min),
                'height': int(y_max - y_min),
                'confidence': float(text_info[1]),
            })
    
    return results


def ocr_image_easyocr(image_path: str) -> List[Dict]:
    """使用EasyOCR识别图片"""
    import easyocr
    
    reader = easyocr.Reader(['ch_sim', 'en'], gpu=False)
    result = reader.readtext(image_path)
    
    results = []
    for detection in result:
        box = detection[0]  # [[x1,y1],[x2,y2],[x3,y3],[x4,y4]]
        text = detection[1]
        confidence = detection[2]
        
        # 计算边界框
        x_coords = [p[0] for p in box]
        y_coords = [p[1] for p in box]
        x_min, x_max = min(x_coords), max(x_coords)
        y_min, y_max = min(y_coords), max(y_coords)
        
        results.append({
            'text': text,
            'x': int(x_min),
            'y': int(y_min),
            'width': int(x_max - x_min),
            'height': int(y_max - y_min),
            'confidence': float(confidence),
        })
    
    return results


def ocr_image(image_path: str, engine: str = None) -> List[Dict]:
    """OCR识别图片，返回文字和位置信息"""
    if not OCR_AVAILABLE:
        raise RuntimeError("没有可用的OCR引擎，请安装 pytesseract 或 paddleocr")
    
    engine = engine or OCR_ENGINE
    
    if engine == 'tesseract':
        return ocr_image_tesseract(image_path)
    elif engine == 'paddleocr':
        return ocr_image_paddleocr(image_path)
    elif engine == 'easyocr':
        return ocr_image_easyocr(image_path)
    else:
        # 默认尝试tesseract
        return ocr_image_tesseract(image_path)


def detect_table_structure(ocr_results: List[Dict], image_width: int, image_height: int) -> List[List[str]]:
    """尝试从OCR结果中检测表格结构"""
    if not ocr_results:
        return []
    
    # 按Y坐标分组（同一行的文字）
    lines = {}
    y_tolerance = 15  # Y坐标容差
    
    for item in ocr_results:
        y = item['y']
        # 找到最近的行
        for line_y in lines:
            if abs(y - line_y) <= y_tolerance:
                lines[line_y].append(item)
                break
        else:
            lines[y] = [item]
    
    # 按Y坐标排序
    sorted_lines = sorted(lines.items(), key=lambda x: x[0])
    
    # 在每行内按X坐标排序
    rows = []
    for line_y, items in sorted_lines:
        sorted_items = sorted(items, key=lambda x: x['x'])
        row_texts = [item['text'] for item in sorted_items]
        rows.append(row_texts)
    
    return rows


def create_excel_from_table(rows: List[List[str]], output_path: str) -> bool:
    """从表格数据创建Excel文件"""
    try:
        import openpyxl
    except ImportError:
        print("[Excel] openpyxl 未安装，尝试使用 pandas")
        try:
            import pandas as pd
            # 使用pandas创建Excel
            df = pd.DataFrame(rows)
            df.to_excel(output_path, index=False, header=False)
            return True
        except ImportError:
            print("[Excel] pandas 也未安装")
            return False
    
    wb = openpyxl.Workbook()
    ws = wb.active
    
    for row_idx, row_data in enumerate(rows, start=1):
        for col_idx, cell_value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=cell_value)
    
    wb.save(output_path)
    return True


def create_word_from_text(texts: List[str], output_path: str) -> bool:
    """从文字列表创建Word文件"""
    try:
        from docx import Document
    except ImportError:
        print("[Word] python-docx 未安装")
        return False
    
    doc = Document()
    
    for text in texts:
        if text.strip():
            doc.add_paragraph(text)
    
    doc.save(output_path)
    return True


def create_word_from_ocr_results(ocr_results: List[Dict], output_path: str) -> bool:
    """从OCR结果创建Word文件，保留大致布局"""
    try:
        from docx import Document
    except ImportError:
        print("[Word] python-docx 未安装")
        return False
    
    doc = Document()
    
    # 按Y坐标分组
    lines = {}
    for item in ocr_results:
        y = item['y']
        if y not in lines:
            lines[y] = []
        lines[y].append(item)
    
    # 按Y排序，每行内按X排序
    sorted_y = sorted(lines.keys())
    
    for y in sorted_y:
        items = sorted(lines[y], key=lambda x: x['x'])
        line_text = ' '.join([item['text'] for item in items])
        if line_text.strip():
            doc.add_paragraph(line_text)
    
    doc.save(output_path)
    return True


async def image_to_excel_async(image_paths: List[str], output_path: str, engine: str = None) -> tuple:
    """异步处理图片转Excel"""
    try:
        all_rows = []
        
        for image_path in image_paths:
            # OCR识别
            ocr_results = await asyncio.to_thread(ocr_image, image_path, engine)
            
            # 获取图片尺寸
            img = Image.open(image_path)
            width, height = img.size
            
            # 检测表格结构
            rows = detect_table_structure(ocr_results, width, height)
            all_rows.extend(rows)
        
        # 创建Excel
        success = create_excel_from_table(all_rows, output_path)
        
        if success:
            return True, output_path, "转换成功"
        else:
            return False, None, "Excel创建失败，请安装 openpyxl 或 pandas"
    
    except Exception as e:
        return False, None, str(e)


async def image_to_word_async(image_paths: List[str], output_path: str, engine: str = None) -> tuple:
    """异步处理图片转Word"""
    try:
        all_texts = []
        all_results = []
        
        for image_path in image_paths:
            # OCR识别
            ocr_results = await asyncio.to_thread(ocr_image, image_path, engine)
            all_results.extend(ocr_results)
            
            # 收集文字
            texts = [item['text'] for item in ocr_results]
            all_texts.extend(texts)
        
        # 创建Word（保留布局）
        success = create_word_from_ocr_results(all_results, output_path)
        
        if success:
            return True, output_path, "转换成功"
        else:
            return False, None, "Word创建失败，请安装 python-docx"
    
    except Exception as e:
        return False, None, str(e)


def get_ocr_status() -> Dict[str, Any]:
    """获取OCR服务状态"""
    return {
        "available": OCR_AVAILABLE,
        "engine": OCR_ENGINE if OCR_AVAILABLE else "none",
        "supportedEngines": ["tesseract", "paddleocr", "easyocr"],
    }